"""Write the keyword map workbook from two coverage runs (scripts/keyword-coverage.js --json).
Usage: python scripts/keyword-map-xlsx.py before.json after.json out.xlsx
"""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

before_f, after_f, out = sys.argv[1:4]
before = {(r['cat'], r['n']): r for r in json.load(open(before_f, encoding='utf8'))}
after = json.load(open(after_f, encoding='utf8'))
CATS = {1: 'Range Rover prices & buying intent', 2: 'Toyota Land Cruiser V8', 3: 'Mercedes-Benz G-Wagon',
        4: 'Car imports from UK to Kenya', 5: 'Car imports from Japan to Kenya', 6: 'Harrier, CX-5 & Prado TX'}
SITE = 'https://elisamotors.co.ke'
# Keywords deliberately not targeted: they are about other businesses or services Elisa Motors does not offer
RULES = [
    ('Not targeted: other company / marketplace name', r'\b(jiji|pigiame|carmudi|dubaicars|sbt|be forward|cmc|inchcape|toyotsu|dt dobie|mhh|autoclassic|kcb|equity bank|hoegh|eukor|carimportskenya|strongero|braneew|car hantei|range rover house|auto terminal)\b'),
    ('Not targeted: car hire / rental (not an Elisa Motors service)', r'\b(hire|rent|rental|wedding)\b'),
    ('Not targeted: repairs, parts & accessories (not an Elisa Motors service)', r'\b(mechanic|garage|repair|spare parts|tires|tyres|tire|exhaust|injector|windshield|android screen|body kit|facelift kit|bullbar|snorkel|alarm|speaker|restoration|actuator|side steps|mirror guard|bush|gear ?box|turbocharger|coolant|carbon buildup|parking brake|dashboard crack|engine price|lumbar|upgrade|leather seats)\b'),
]


def note(kw, level):
    for label, rx in RULES:
        if re.search(rx, kw):
            return label
    return {'strong': 'Targeted in a title, heading or FAQ', 'weak': 'Covered in page content'}.get(level, 'Gap: no page yet (candidate for a future guide)')


F = lambda **k: Font(name='Arial', **k)
HDR = PatternFill('solid', start_color='0E2A3A')
LINE = Border(bottom=Side(style='thin', color='D9D9D9'))
LV = {'strong': 'Strong', 'weak': 'Weak', 'missing': 'Missing'}

wb = Workbook()
ws = wb.active
ws.title = 'Keyword Map'
ws.append(['#', 'Category', 'Keyword', 'Target page', 'Coverage before', 'Coverage after', 'Status / note'])
for r in after:
    b = before[(r['cat'], r['n'])]
    page = (r['pages'] or b['pages'] or [''])[0]
    ws.append([r['n'], CATS[r['cat']], r['kw'], SITE + page if page else '', LV[b['level']], LV[r['level']], note(r['kw'], r['level'])])
for c in ws[1]:
    c.font, c.fill = F(bold=True, color='FFFFFF'), HDR
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font, c.border = F(size=10), LINE
    if row[3].value:
        row[3].hyperlink = row[3].value
        row[3].font = F(size=10, color='0563C1', underline='single')
for i, w in enumerate([6, 32, 52, 62, 16, 16, 60], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:G{ws.max_row}'
n = ws.max_row

s = wb.create_sheet('Summary', 0)
s['A1'] = 'Elisa Motors: keyword coverage (elisamotors.co.ke)'
s['A1'].font = F(bold=True, size=14)
s['A2'] = 'Source: "elisa motors research.xlsx" (951 keywords, provided by Elisa Motors). Coverage measured on the built site before and after the October 2026 SEO update.'
s['A2'].font = F(italic=True, size=9, color='666666')
s.append([])
s.append(['Category', 'Keywords', 'Strong before', 'Strong after', 'Missing before', 'Missing after', 'Not targeted on purpose'])
for c in s[4]:
    c.font, c.fill = F(bold=True, color='FFFFFF'), HDR
rng = lambda col: f"'Keyword Map'!${col}$2:${col}${n}"
for i, name in enumerate(CATS.values()):
    r = 5 + i
    s.append([name, f'=COUNTIFS({rng("B")},A{r})',
              f'=COUNTIFS({rng("B")},A{r},{rng("E")},"Strong")', f'=COUNTIFS({rng("B")},A{r},{rng("F")},"Strong")',
              f'=COUNTIFS({rng("B")},A{r},{rng("E")},"Missing")', f'=COUNTIFS({rng("B")},A{r},{rng("F")},"Missing")',
              f'=COUNTIFS({rng("B")},A{r},{rng("G")},"Not targeted*")'])
t = 5 + len(CATS)
s.append(['Total'] + [f'=SUM({get_column_letter(c)}5:{get_column_letter(c)}{t - 1})' for c in range(2, 8)])
for row in s.iter_rows(min_row=5, max_row=t):
    for c in row:
        c.font, c.border = F(size=10, bold=(c.row == t)), LINE
s.append([])
for k, v in [('Strong', "All the keyword's main words appear in one page title, heading or FAQ question."),
             ('Weak', "The keyword's words appear in a page's content, but no heading targets it."),
             ('Missing', 'No page on the site covers it.'),
             ('Not targeted on purpose', "Other companies' names, car hire, and repairs/parts/accessories. Elisa Motors does not offer these, so pages for them would mislead visitors.")]:
    s.append([k, v])
    s.cell(s.max_row, 1).font, s.cell(s.max_row, 2).font = F(bold=True, size=10), F(size=10)
s.column_dimensions['A'].width = 34
for col in 'BCDEFG':
    s.column_dimensions[col].width = 17
wb.calculation.fullCalcOnLoad = True  # Excel computes the summary formulas on open
wb.save(out)
print('saved', out, n - 1, 'keywords')
